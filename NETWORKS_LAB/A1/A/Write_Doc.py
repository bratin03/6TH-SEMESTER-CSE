# Import necessary modules
from scapy.all import *
import re
from docx import Document

# Function to process DNS packets and update a dictionary with Q and aggregated A values
def process_dns(pkt, q_a_dict):
    if DNS in pkt and pkt.haslayer(DNSQR):
        query = pkt[DNSQR].qname.decode('utf-8')
        responses = []

        # Check if DNS response records are present
        if DNSRR in pkt:
            dnsrr_strings = repr(pkt[DNSRR]).split("|")
            rrname_pattern = re.compile(r"rrname='(.*?)'")
            rdata_pattern = re.compile(r"rdata=(.*?) ")
            a_values = []

            # Extract RData values
            for entry in dnsrr_strings:
                rdata_match = rdata_pattern.search(entry)
                if rdata_match:
                    rdata = rdata_match.group(1)
                    a_values.append(rdata)

            # Update dictionary with Q and A values
            if query in q_a_dict:
                q_a_dict[query].extend(a_values)
            else:
                q_a_dict[query] = a_values

# Main function
def main():
    # Specify the pcap file name
    pcap_file = '21CS10016_CAPTURE.pcap'
    
    # Create a Document object for creating a Word document
    doc = Document()

    # Dictionary to store Q and aggregated A values
    q_a_dict = {}

    try:
        # Read pcap file using Scapy
        pkts = rdpcap(pcap_file)

        # Process each packet in the pcap file
        for pkt in pkts:
            process_dns(pkt, q_a_dict)

        # Create a single table with header row Q and A
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = 200  # Adjust column widths as needed

        # Add headers to the table
        table.cell(0, 0).text = 'Q'
        table.cell(0, 1).text = 'A'

        # Populate the table with Q and aggregated A values
        for q, a_values in q_a_dict.items():
            row_cells = table.add_row().cells
            row_cells[0].text = q
            row_cells[1].text = ', '.join(a_values)

        # Save the document to a file
        doc.save('dns_results.docx')
        print("Results saved to dns_results.docx")
    
    # Handle file not found exception
    except FileNotFoundError:
        print(f"Error: File {pcap_file} not found.")
    
    # Handle Scapy exception
    except Scapy_Exception as e:
        print(f"Error: {e}")

# Run the main function if the script is executed
if __name__ == "__main__":
    main()

